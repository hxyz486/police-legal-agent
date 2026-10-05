# -*- coding: utf-8 -*-
"""读取输入 Excel：Sheet「警情数据」，字段：反馈单编号 / 出警情况。"""
import logging
from openpyxl import load_workbook

from . import config

log = logging.getLogger("saiti.reader")

HEADER_ID = "反馈单编号"
HEADER_TEXT = "出警情况"


def read_records(path: str = None):
    """返回 [(反馈单编号, 出警情况), ...]，保持输入顺序。

    容错：Sheet 名找不到取第一个 Sheet；列名不匹配按列位置兜底；
    空行/编号缺失的记录跳过并记日志，绝不让单条脏数据拖垮整体。
    """
    path = path or config.INPUT_PATH
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["警情数据"] if "警情数据" in wb.sheetnames else wb[wb.sheetnames[0]]
        if "警情数据" not in wb.sheetnames:
            log.warning("未找到 Sheet「警情数据」，改用第一个 Sheet：%s", ws.title)
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    if not rows:
        log.error("输入文件为空：%s", path)
        return []

    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    if HEADER_ID in header and HEADER_TEXT in header:
        id_col = header.index(HEADER_ID)
        text_col = header.index(HEADER_TEXT)
    else:
        log.warning("表头与规范不符(%s)，按第 1、2 列兜底", header)
        id_col, text_col = 0, 1

    records = []
    for i, row in enumerate(rows[1:], start=2):
        def cell(idx):
            return row[idx] if idx < len(row) else None

        rid = cell(id_col)
        text = cell(text_col)
        if rid is None and (text is None or str(text).strip() == ""):
            log.warning("第 %d 行为空行，跳过", i)
            continue
        if rid is None:
            log.warning("第 %d 行反馈单编号缺失，跳过", i)
            continue
        records.append((str(rid).strip(), "" if text is None else str(text)))
    log.info("读取记录 %d 条：%s", len(records), path)
    return records
