# -*- coding: utf-8 -*-
"""输出格式合规自检：Sheet 名/表头/行数/编号顺序/枚举值逐字比对规范。"""
import logging

from openpyxl import load_workbook

from . import writer

log = logging.getLogger("saiti.validate")


def validate(path: str, input_records) -> list:
    """对照输入记录校验输出文件。返回问题列表（空 = 通过）。"""
    issues = []
    wb = load_workbook(path, read_only=False, data_only=True)
    try:
        if wb.sheetnames != ["风险研判结果"]:
            issues.append(f"Sheet 名不符：{wb.sheetnames}")
            return issues
        ws = wb["风险研判结果"]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    if not rows:
        return ["输出为空"]
    header = [str(c) if c is not None else "" for c in rows[0]]
    if header != writer.HEADERS:
        issues.append(f"表头不符：{header}")

    body = rows[1:]
    if len(body) != len(input_records):
        issues.append(f"行数不符：输出 {len(body)} != 输入 {len(input_records)}")

    for i, ((rid_in, _), row) in enumerate(zip(input_records, body), start=2):
        if row is None or len(row) < 7:
            issues.append(f"第 {i} 行列数不足")
            continue
        rid_out, _, names, ids, exists, level, _ = row[:7]
        names = "" if names is None else str(names)
        ids = "" if ids is None else str(ids)
        # 官方样例集为布尔值，程序输出兼容布尔与 true/false 字符串两种形态
        ex_norm = "true" if exists in (True, "true", "True", "TRUE") else \
            "false" if exists in (False, "false", "False", "FALSE") else str(exists)
        if str(rid_out) != str(rid_in):
            issues.append(f"第 {i} 行编号不对应：{rid_out} != {rid_in}")
        if ex_norm not in ("true", "false"):
            issues.append(f"第 {i} 行是否存在风险人员非法：{exists}")
        if str(level) not in ("高", "低", "无"):
            issues.append(f"第 {i} 行风险等级非法：{level}")
        if ex_norm == "false" and (names.strip() or ids.strip()):
            issues.append(f"第 {i} 行标记无风险但姓名/证号非空")
        if ex_norm == "true" and str(level) == "无":
            issues.append(f"第 {i} 行存在风险人员但等级为无")
    return issues
