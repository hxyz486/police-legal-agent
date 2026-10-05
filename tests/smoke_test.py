# -*- coding: utf-8 -*-
"""端到端冒烟测试：mock 模型服务 + 真实服务进程 + 批量流水线。

覆盖：
1. mock 模型服务（LLM/Embedding/Rerank）下启动统一服务；
2. GET /health、POST /qa（RAG 问答）、POST /assess（法条接地研判）；
3. 批量研判 CLI（assess_batch.py）输入 xlsx -> 输出 xlsx。

运行：python tests/smoke_test.py
"""
import json
import os
import subprocess
import sys
import time
import urllib.request

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)

import mock_model_server  # noqa: E402

PORT = 18990
BASE = f"http://127.0.0.1:{PORT}"


def http(method, path, body=None, timeout=120):
    req = urllib.request.Request(
        BASE + path,
        data=json.dumps(body).encode("utf-8") if body is not None else None,
        headers={"Content-Type": "application/json"},
        method=method)
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return json.loads(resp.read().decode("utf-8"))


def wait_health(proc, deadline=180):
    t0 = time.time()
    while time.time() - t0 < deadline:
        if proc.poll() is not None:
            raise RuntimeError("服务进程提前退出，请检查日志输出")
        try:
            h = http("GET", "/health", timeout=5)
            return h
        except Exception:  # noqa: BLE001
            time.sleep(2)
    raise RuntimeError("服务未在时限内就绪")


def main():
    # 1. mock 模型服务
    import threading
    from socketserver import ThreadingMixIn
    from http.server import HTTPServer

    class Srv(ThreadingMixIn, HTTPServer):
        daemon_threads = True

    mock = Srv(("127.0.0.1", mock_model_server.PORT), mock_model_server.Mock)
    threading.Thread(target=mock.serve_forever, daemon=True).start()
    print("[OK] mock 模型服务 :", mock_model_server.PORT)

    env = dict(os.environ)
    mu = f"http://127.0.0.1:{mock_model_server.PORT}"
    env.update({
        "LLM_API_URL": mu,
        "EMBEDDING_API_URL": mu,
        "RERANK_API_URL": mu,
        "LLM_MODEL": "mock-llm",
        "LAWS_DIR": os.path.join(ROOT, "laws"),
        "LOG_PATH": os.path.join(ROOT, "log", "smoke_service.log"),
        "CACHE_PATH": os.path.join(ROOT, "data", "smoke_cache.json"),
        "PRESET_FILE": os.path.join(ROOT, "preset", "preset_answers.json"),
        "PORT": str(PORT),
        "PYTHONIOENCODING": "utf-8",
    })

    # 2. 启动统一服务（stdout 重定向文件：PIPE 无人读取会写满缓冲区卡死服务）
    svc_log = open(os.path.join(ROOT, "log", "smoke_stdout.log"), "w", encoding="utf-8")
    proc = subprocess.Popen([sys.executable, os.path.join(ROOT, "src", "service.py")],
                            env=env, cwd=ROOT, stdout=svc_log, stderr=subprocess.STDOUT)
    try:
        h = wait_health(proc)
        print("[OK] /health:", h)

        # 3. /qa 法律问答
        qa = http("POST", "/qa", {"question": "民警在巡逻中发现有人赌博，应如何处理？"})
        assert isinstance(qa.get("answer"), str) and qa["answer"], qa
        assert isinstance(qa.get("sources"), list), qa
        for s in qa["sources"]:
            assert set(s.keys()) == {"law_name", "article", "snippet"}, s
        print(f"[OK] /qa: answer={qa['answer'][:40]}... sources={len(qa['sources'])}")

        # 4. /assess 警情研判（含法条接地）
        text = ("【当事人信息】1.报警人：王XX，男，1985年XX月，331010XXXXXXXX0416；"
                "2.当事人：赵XX，男，1990年XX月，321120XXXXXXXX0416。"
                "【警情内容及处置情况】赵XX因债务纠纷扬言\"再不还钱就砍死你全家\"，"
                "并持菜刀到王XX店门口叫嚣，民警到场将其控制。")
        ar = http("POST", "/assess", {"text": text})
        for k in ("exists", "level", "risk_persons", "law_references", "degraded", "reason"):
            assert k in ar, (k, ar)
        for p in ar["risk_persons"]:
            assert set(p.keys()) == {"name", "id_number", "level", "reason"}, p
        for r in ar["law_references"]:
            assert set(r.keys()) == {"law_name", "article", "snippet"}, r
        # mock 模型对持械威胁样例返回确定性研判：必须识别出高风险人员
        assert ar["exists"] is True and ar["level"] == "高", ar
        assert len(ar["risk_persons"]) == 1 and ar["risk_persons"][0]["name"] == "赵XX", ar
        # 法条引用只列研判理由实际援引的条文（truthful refs）：结构必须是三键
        assert isinstance(ar["law_references"], list), ar
        for r in ar["law_references"]:
            assert set(r.keys()) == {"law_name", "article", "snippet"}, r
        assert ar["degraded"] is False, ar
        print(f"[OK] /assess: exists={ar['exists']} level={ar['level']} "
              f"persons={len(ar['risk_persons'])} refs={len(ar['law_references'])} "
              f"degraded={ar['degraded']}")

        # 5. 批量研判 CLI
        from openpyxl import Workbook, load_workbook
        in_x = os.path.join(HERE, "smoke_input.xlsx")
        out_x = os.path.join(HERE, "smoke_output.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "警情数据"
        ws.append(["反馈单编号", "出警情况"])
        ws.append(["JQ-001", text])
        ws.append(["JQ-002", "报警人称电动车被盗，情绪平稳，无过激言行。"])
        wb.save(in_x)

        benv = dict(env)
        benv["LLM_MODEL"] = "mock-llm"
        r = subprocess.run(
            [sys.executable, os.path.join(ROOT, "src", "assess_batch.py"),
             in_x, out_x],
            env=benv, cwd=ROOT, capture_output=True, text=True,
            encoding="utf-8", errors="replace", timeout=600)
        if r.returncode != 0:
            print(r.stdout[-3000:])
            print(r.stderr[-3000:])
            raise RuntimeError("批量研判退出码非 0")
        wb2 = load_workbook(out_x)
        assert wb2.sheetnames == ["风险研判结果"], wb2.sheetnames
        rows = list(wb2["风险研判结果"].iter_rows(values_only=True))
        assert rows[0][0] == "反馈单编号" and len(rows) == 3, rows
        print(f"[OK] 批量: 输出 {len(rows)-1} 行，表头/结构合规")
    finally:
        proc.terminate()
        try:
            proc.wait(timeout=10)
        except Exception:  # noqa: BLE001
            proc.kill()
        svc_log.close()
        mock.shutdown()

    print("\nSMOKE TEST PASSED")


if __name__ == "__main__":
    main()
